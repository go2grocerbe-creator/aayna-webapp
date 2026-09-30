#!/usr/bin/env python3
"""
AAYNA real-inventory importer.

    AAYNA_Product_Inventory.xlsx + Product_Photos_Higgsfield/
        -> deterministic filename matching
        -> validated dry-run manifest (JSON)
        -> existing products collection + existing object storage, as DRAFT

Two steps, deliberately separate:

  1) manifest (dry run, touches nothing):
     python scripts/import_inventory.py manifest \
         --xlsx "E:/deMarkt/AAYNA Brand Details/AAYNA_Product_Inventory.xlsx" \
         --photos "E:/deMarkt/AAYNA - Product Sourcing Photos/Product_Photos_Higgsfield" \
         --out import_manifest.json

  2) import (writes to the DB configured in backend/.env):
     python scripts/import_inventory.py import --manifest import_manifest.json \
         [--only AYN-RNG-CLS-GLD] [--yes]

Rules this script enforces (see task brief):
  * The spreadsheet is the source of truth. Nothing is invented: no selling
    price, discount, material, description or category is guessed. A product
    without a selling price is NOT ready and is skipped.
  * Selling price comes from a "Selling price (BDT)" column if present, else
    from --temp-price (founder-approved temporary price, flagged
    price_is_temporary). Final prices are managed in Admin > Products.
  * Images go through the app's existing object storage (storage.put_object +
    db.files record), exactly like the admin upload endpoint. Local Windows
    paths never reach a product record. Originals are never modified; an
    optimised WebP copy (max 1600px) is uploaded instead of 2-8 MB PNGs.
  * Idempotent: products are matched by SKU. Re-running only refreshes
    spreadsheet-owned fields (name, category, colour, size, cost, notes).
    Selling/discount price, status, images and stock are owned by Admin once
    the product exists and are never overwritten. Identical image bytes are
    never re-uploaded (sha256 on db.files).
  * New products are created with status "draft", which the public API
    already hides (it only serves "active" / "out_of_stock").
  * Unit cost -> cost_price and Notes -> internal_notes: both admin-only. The
    public API uses an explicit whitelist (PUBLIC_PRODUCT_FIELDS) that
    excludes them.
"""
import argparse
import asyncio
import hashlib
import io
import json
import re
import sys
import uuid
from collections import Counter, defaultdict
from pathlib import Path

IMAGE_EXT = {".png", ".jpg", ".jpeg", ".webp"}
MAX_EDGE = 1600
WEBP_QUALITY = 85

# Spreadsheet header -> website field. Headers are matched case-insensitively.
COLUMN_MAP = {
    "sku": "sku",
    "product name": "product_name",
    "category": "_category",            # mapped via CATEGORY_MAP below
    "color / finish": "color",
    "size specification": "size",
    "stock unit": "_stock_unit",        # informational (Piece / Pair)
    "qty": "stock_quantity",            # sellable units (pairs for earrings)
    "individual items": None,           # not a website field
    "unit cost (bdt)": "cost_price",    # ADMIN ONLY
    "unit cost (rmb)": None,            # derivable, not stored
    "stock value (bdt)": None,
    "stock value (rmb)": None,
    "notes": "internal_notes",          # ADMIN ONLY
    # Optional founder-supplied columns (not in the sheet yet):
    "selling price (bdt)": "selling_price",
    "discount price (bdt)": "discount_price",
    "material": "material",
}

# Spreadsheet category -> existing website category slug.
# (value, needs_confirmation)
CATEGORY_MAP = {
    "necklace / locket": ("necklaces", False),
    "earrings": ("earrings", False),
    "bracelet": ("bracelets", False),
    "ring": ("rings", False),
    "cuff": ("bracelets", True),  # no "Cuffs" category exists on the site
}

STOPWORDS = {"aayna", "statement"}


def _norm_tokens(text):
    toks = re.split(r"[^a-z0-9]+", (text or "").lower())
    out = set()
    for t in toks:
        if not t or t in STOPWORDS:
            continue
        if len(t) > 3 and t.endswith("s") and not t.endswith("ss"):
            t = t[:-1]  # earrings -> earring
        out.add(t)
    return out


# ---------------------------------------------------------------------------
# Spreadsheet
# ---------------------------------------------------------------------------
def read_inventory(xlsx_path):
    import openpyxl

    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws = wb["Inventory"] if "Inventory" in wb.sheetnames else wb.active
    rows = list(ws.iter_rows(values_only=True))
    products, header, in_rejected = [], None, False
    for idx, row in enumerate(rows, start=1):
        first = (str(row[0]).strip().lower() if row and row[0] is not None else "")
        if first.startswith("rejected"):
            in_rejected = True
            continue
        if first == "sku":
            header = [(str(c).strip().lower() if c is not None else "") for c in row]
            continue
        if header is None or in_rejected or not any(c is not None for c in row):
            continue
        rec = {"_row": idx}
        for h, v in zip(header, row):
            if h:
                rec[h] = v
        products.append(rec)
    return {"sheet": ws.title, "header": header or [], "products": products}


def map_row(raw, temp_price=None):
    mapped, warnings, blockers = {}, [], []
    for h, v in raw.items():
        if h == "_row":
            continue
        target = COLUMN_MAP.get(h, "__unknown__")
        if target == "__unknown__":
            warnings.append(f"Unmapped column '{h}' ignored")
            continue
        if target is None:
            continue
        mapped[target] = v.strip() if isinstance(v, str) else v

    sku = mapped.get("sku")
    if not sku:
        blockers.append("Missing SKU")
    if not mapped.get("product_name"):
        blockers.append("Missing product name")

    cat_raw = (mapped.pop("_category", "") or "").strip()
    cat = CATEGORY_MAP.get(cat_raw.lower())
    if not cat:
        blockers.append(f"Category '{cat_raw}' has no website category")
    else:
        mapped["category_slug"] = cat[0]
        if cat[1]:
            warnings.append(f"Category '{cat_raw}' mapped to '{cat[0]}' - founder to confirm")

    qty = mapped.get("stock_quantity")
    if not isinstance(qty, (int, float)) or qty < 0 or int(qty) != qty:
        blockers.append(f"Invalid stock quantity: {qty!r}")
    else:
        mapped["stock_quantity"] = int(qty)

    for f in ("cost_price", "selling_price", "discount_price"):
        v = mapped.get(f)
        if v in (None, ""):
            mapped.pop(f, None)
            continue
        if not isinstance(v, (int, float)) or v < 0:
            blockers.append(f"Invalid {f}: {v!r}")
    if "selling_price" not in mapped and temp_price is not None:
        # Founder-approved TEMPORARY price; final price is set in Admin.
        # Never written back to the spreadsheet.
        mapped["selling_price"] = float(temp_price)
        mapped["price_is_temporary"] = True
        warnings.append(f"Temporary price {temp_price:g} BDT - set final price in Admin")
    if "selling_price" not in mapped:
        blockers.append("Missing selling price (BDT) - not in spreadsheet")
    elif mapped["selling_price"] <= 0:
        blockers.append("Selling price must be > 0")
    dp = mapped.get("discount_price")
    if dp is not None and "selling_price" in mapped and dp >= mapped["selling_price"]:
        blockers.append("Discount price must be lower than selling price")

    # SKU colour code vs colour column
    color = (mapped.get("color") or "").lower()
    code = (sku or "").rsplit("-", 1)[-1].upper()
    code_color = {"GLD": "golden", "SLV": "silver"}.get(code)
    if code_color and color and code_color != color:
        warnings.append(f"SKU suffix {code} says {code_color} but Color column says {color}")
    return mapped, warnings, blockers


# ---------------------------------------------------------------------------
# Photos
# ---------------------------------------------------------------------------
def index_photos(photo_dir):
    files = []
    root = Path(photo_dir)
    for p in sorted(root.rglob("*")):
        if p.is_file():
            stem = p.stem
            m = re.search(r"[\s_-]+(\d+)$", stem)
            files.append({
                "filename": p.name,
                "relative_path": str(p.relative_to(root)).replace("\\", "/"),
                "extension": p.suffix.lower(),
                "size": p.stat().st_size,
                "order": int(m.group(1)) if m else None,
                "tokens": sorted(_norm_tokens(stem[: m.start()] if m else stem)),
            })
    return files


def match_photos(products, photos):
    """A photo belongs to a product when every distinguishing token of the
    product (name + colour, minus 'Aayna'/'Statement') appears in the photo
    filename. Word order is ignored. Exactly one candidate -> MATCHED; more
    than one -> AMBIGUOUS; a subset match without colour/number -> PARTIAL."""
    ptoks = {p["sku"]: _norm_tokens(f"{p['product_name']} {p.get('color', '')}") for p in products}
    assign, partial, orphans, ambiguous = defaultdict(list), defaultdict(list), [], []
    for ph in photos:
        if ph["extension"] not in IMAGE_EXT:
            orphans.append({**ph, "reason": "unsupported format"})
            continue
        ftoks = set(ph["tokens"])
        full = [sku for sku, t in ptoks.items() if t and t <= ftoks]
        if len(full) == 1 and ph["order"] is not None:
            assign[full[0]].append(ph)
            continue
        if len(full) > 1:
            ambiguous.append({**ph, "candidates": full})
            continue
        # partial: filename lacks colour / sequence number but otherwise fits
        near = [sku for sku, t in ptoks.items()
                if len(t & ftoks) >= max(2, len(t) - 1)]
        if len(near) == 1 or full:
            partial[(full or near)[0]].append(ph)
        elif len(near) > 1:
            ambiguous.append({**ph, "candidates": near})
        else:
            orphans.append({**ph, "reason": "no inventory product"})
    for sku in assign:
        assign[sku].sort(key=lambda x: (x["order"], x["filename"]))
    return assign, partial, orphans, ambiguous


# ---------------------------------------------------------------------------
# Manifest
# ---------------------------------------------------------------------------
def build_manifest(xlsx, photo_dir, temp_price=None):
    inv = read_inventory(xlsx)
    records = []
    for raw in inv["products"]:
        mapped, warnings, blockers = map_row(raw, temp_price)
        records.append({"source_row": raw["_row"], **mapped, "_w": warnings, "_b": blockers})

    # duplicates
    sku_count = Counter(r.get("sku") for r in records)
    name_count = Counter((r.get("product_name") or "").lower() for r in records)
    for r in records:
        if sku_count[r.get("sku")] > 1:
            r["_b"].append("Duplicate SKU in spreadsheet")
        if name_count[(r.get("product_name") or "").lower()] > 1:
            # same name, different colour -> disambiguate display name from the
            # spreadsheet's own Color column (no invented text)
            r["display_name"] = f"{r['product_name']} - {str(r.get('color', '')).title()}"
            r["_w"].append("Name shared with another colour; colour appended to display name/slug")
        else:
            r["display_name"] = r.get("product_name")

    photos = index_photos(photo_dir)
    assign, partial, orphans, ambiguous = match_photos(records, photos)
    seen = Counter(ph["relative_path"] for v in assign.values() for ph in v)

    out = []
    for r in records:
        sku = r.get("sku")
        imgs = assign.get(sku, [])
        orders = [i["order"] for i in imgs]
        if len(set(orders)) != len(orders):
            r["_b"].append("Duplicate image sequence numbers")
        if any(seen[i["relative_path"]] > 1 for i in imgs):
            r["_b"].append("Image assigned to more than one product")
        if not imgs:
            r["_b"].append("No matched images")
        if partial.get(sku):
            r["_w"].append("Extra unnumbered image(s) not auto-attached: "
                           + ", ".join(p["filename"] for p in partial[sku]))
        if imgs:
            match = "MATCHED"
        elif partial.get(sku):
            match = "PARTIAL"
        else:
            match = "NO IMAGE"
        rec = {k: v for k, v in r.items() if not k.startswith("_")}
        rec.update({
            "image_match": match,
            "images": [i["relative_path"] for i in imgs],
            "image_count": len(imgs),
            "hero_image": imgs[0]["relative_path"] if imgs else None,
            "warnings": r["_w"],
            "blockers": r["_b"],
            "ready_to_import": not r["_b"],
        })
        out.append(rec)

    summary = {
        "spreadsheet": str(xlsx), "sheet": inv["sheet"], "photo_dir": str(photo_dir),
        "products": len(out),
        "ready": sum(r["ready_to_import"] for r in out),
        "blocked": sum(not r["ready_to_import"] for r in out),
        "image_match": dict(Counter(r["image_match"] for r in out)),
        "photo_files": len(photos),
        "ambiguous_images": [{"file": a["relative_path"], "candidates": a["candidates"]} for a in ambiguous],
        "partial_images": [p["relative_path"] for v in partial.values() for p in v],
        "orphan_images": [{"file": o["relative_path"], "reason": o["reason"]} for o in orphans],
        "unmapped_headers": sorted(h for h in inv["header"] if h and h not in COLUMN_MAP),
        "blocker_counts": dict(Counter(b for r in out for b in r["blockers"])),
    }
    return {"summary": summary, "products": out}


# ---------------------------------------------------------------------------
# Import (uses the backend's own db + storage modules)
# ---------------------------------------------------------------------------
def optimise_image(path):
    from PIL import Image, ImageOps

    im = ImageOps.exif_transpose(Image.open(path))
    im = im.convert("RGB")
    im.thumbnail((MAX_EDGE, MAX_EDGE), Image.LANCZOS)
    buf = io.BytesIO()
    im.save(buf, "WEBP", quality=WEBP_QUALITY, method=6)
    return buf.getvalue()


async def upload_image(db, storage, src_path, public_base):
    raw = Path(src_path).read_bytes()
    sha = hashlib.sha256(raw).hexdigest()
    existing = await db.files.find_one({"source_sha256": sha, "is_deleted": False})
    if existing:
        return existing["storage_path"], False
    data = optimise_image(src_path)
    path = f"{storage.APP_NAME}/uploads/{uuid.uuid4().hex}.webp"
    result = await asyncio.to_thread(storage.put_object, path, data, "image/webp")
    stored = result.get("path", path)
    await db.files.insert_one({
        "id": str(uuid.uuid4()), "storage_path": stored,
        "original_filename": Path(src_path).name, "content_type": "image/webp",
        "size": result.get("size", len(data)), "is_deleted": False,
        "source_sha256": sha, "created_at": storage.now_iso(),
    })
    return stored, True


async def import_products(manifest, photo_dir, public_base, only=None, db=None, storage=None):
    if db is None:
        backend = Path(__file__).resolve().parent.parent / "backend"
        sys.path.insert(0, str(backend))
        from db import db  # noqa: E402
        import storage  # noqa: E402
    from admin_routes import slugify, _log_inventory  # reuse existing helpers

    report = {"created": [], "updated": [], "skipped": [], "failed": [], "images_uploaded": 0}
    for rec in manifest["products"]:
        sku = rec.get("sku")
        if only and sku not in only:
            continue
        if not rec["ready_to_import"]:
            report["skipped"].append({"sku": sku, "blockers": rec["blockers"]})
            continue
        try:
            cat = await db.categories.find_one({"slug": rec["category_slug"]}, {"_id": 0})
            if not cat:
                raise RuntimeError(f"category '{rec['category_slug']}' missing in DB")
            # Spreadsheet-owned fields: refreshed on every run.
            fields = {
                "product_name": rec["display_name"], "sku": sku,
                "category_slug": cat["slug"], "category_name": cat["name"],
                "cost_price": float(rec.get("cost_price") or 0),
                "color": rec.get("color") or "", "size": rec.get("size") or "",
                "internal_notes": rec.get("internal_notes") or "",
                "source": {"type": "inventory_xlsx", "row": rec["source_row"]},
                "updated_at": storage.now_iso(),
            }
            if rec.get("material"):
                fields["material"] = rec["material"]
            existing = await db.products.find_one({"sku": sku})
            if existing:
                # Admin-owned after creation: selling/discount price, status,
                # images and stock are never overwritten by a re-run.
                await db.products.update_one({"sku": sku}, {"$set": fields})
                report["updated"].append(sku)
                continue
            images = []
            for i, rel in enumerate(rec["images"]):
                stored, new = await upload_image(db, storage, Path(photo_dir) / rel, public_base)
                report["images_uploaded"] += int(new)
                images.append({
                    "image_url": f"{public_base.rstrip('/')}/api/files/{stored}",
                    "alt_text": f"{rec['display_name']} - image {i + 1}",
                    "is_main": i == 0, "sort_order": i,
                })
            slug = slugify(rec["display_name"])
            if await db.products.find_one({"slug": slug}):
                slug = f"{slug}-{slugify(sku)}"
            doc = {
                "id": str(uuid.uuid4()), "slug": slug, **fields,
                "selling_price": float(rec["selling_price"]),
                "discount_price": float(rec["discount_price"]) if rec.get("discount_price") else None,
                "price_is_temporary": bool(rec.get("price_is_temporary")),  # admin-only marker
                "images": images,
                "stock_quantity": rec["stock_quantity"], "low_stock_alert": 3,
                "short_description": "", "full_description": "", "material": fields.get("material", ""),
                "weight": None, "status": "draft",
                "is_featured": False, "is_best_seller": False, "is_new_arrival": False,
                "tags": [], "category_id": None, "brand_id": None, "seller_id": None,
                "created_at": storage.now_iso(),
            }
            await db.products.insert_one({**doc})
            if doc["stock_quantity"] > 0:
                await _log_inventory(doc["id"], "stock_in", doc["stock_quantity"], 0,
                                     doc["stock_quantity"], "Inventory spreadsheet import", None)
            report["created"].append(sku)
        except Exception as exc:  # noqa: BLE001
            report["failed"].append({"sku": sku, "error": str(exc)})
    return report


def main():
    ap = argparse.ArgumentParser()
    sub = ap.add_subparsers(dest="cmd", required=True)
    m = sub.add_parser("manifest")
    m.add_argument("--xlsx", required=True)
    m.add_argument("--photos", required=True)
    m.add_argument("--out", default="import_manifest.json")
    m.add_argument("--temp-price", type=float,
                   help="founder-approved temporary selling price (BDT) for rows without one; products stay draft")
    i = sub.add_parser("import")
    i.add_argument("--manifest", required=True)
    i.add_argument("--photos", help="override photo dir from manifest")
    i.add_argument("--public-base", help="public backend URL (default REACT_APP_BACKEND_URL / PUBLIC_BACKEND_URL)")
    i.add_argument("--only", action="append", help="SKU to import (repeatable) - use for the single-product test")
    i.add_argument("--yes", action="store_true")
    a = ap.parse_args()

    if a.cmd == "manifest":
        man = build_manifest(a.xlsx, a.photos, a.temp_price)
        Path(a.out).write_text(json.dumps(man, indent=2, ensure_ascii=False), encoding="utf-8")
        print(json.dumps(man["summary"], indent=2, ensure_ascii=False))
        return

    import os
    man = json.loads(Path(a.manifest).read_text(encoding="utf-8"))
    photo_dir = a.photos or man["summary"]["photo_dir"]
    base = a.public_base or os.environ.get("PUBLIC_BACKEND_URL") or os.environ.get("REACT_APP_BACKEND_URL")
    if not base:
        sys.exit("Set --public-base to the backend's public URL (image URLs are built from it).")
    backend = Path(__file__).resolve().parent.parent / "backend"
    sys.path.insert(0, str(backend))
    import db as dbmod  # noqa: E402
    print(f"Target DB: {os.environ.get('DB_NAME')}  |  images -> {base}/api/files/...")
    if not a.yes and input("Type IMPORT to continue: ").strip() != "IMPORT":
        sys.exit("Aborted.")
    rep = asyncio.run(import_products(man, photo_dir, base, set(a.only) if a.only else None, db=dbmod.db,
                                      storage=__import__("storage")))
    print(json.dumps(rep, indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()
